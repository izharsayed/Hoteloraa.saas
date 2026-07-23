import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Design System Colors
NAVY = RGBColor(11, 31, 58) # #0B1F3A
IVORY = RGBColor(254, 249, 241) # #FEF9F1
GOLD = RGBColor(212, 175, 55) # #D4AF37
SLATE = RGBColor(107, 114, 128) # #6B7280

def add_slide(prs, title, content_items):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    # Background color
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = IVORY
    
    title_shape = slide.shapes.title
    title_shape.text = title
    title_shape.text_frame.paragraphs[0].font.color.rgb = NAVY
    title_shape.text_frame.paragraphs[0].font.name = 'Lora'
    title_shape.text_frame.paragraphs[0].font.bold = True
    
    content_shape = slide.placeholders[1]
    tf = content_shape.text_frame
    tf.text = content_items[0]
    tf.paragraphs[0].font.color.rgb = SLATE
    tf.paragraphs[0].font.name = 'Inter'
    
    for item in content_items[1:]:
        p = tf.add_paragraph()
        p.text = item
        p.font.color.rgb = SLATE
        p.font.name = 'Inter'

def create_presentation():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY
    
    title = slide.shapes.title
    title.text = "Hoteloraa SaaS"
    title.text_frame.paragraphs[0].font.color.rgb = GOLD
    title.text_frame.paragraphs[0].font.name = 'Lora'
    
    subtitle = slide.placeholders[1]
    subtitle.text = "Premium Hospitality Management Platform"
    subtitle.text_frame.paragraphs[0].font.color.rgb = IVORY
    subtitle.text_frame.paragraphs[0].font.name = 'Inter'
    
    # Slide 1: Introduction
    add_slide(prs, "What is Hoteloraa?", [
        "A multi-tenant SaaS platform tailored for the hospitality industry.",
        "Supports three main business types:",
        "  - Restaurant Management",
        "  - Lodging / Hotel Management",
        "  - Hybrid (Hotel + Restaurant)",
        "Premium UI designed for operational efficiency."
    ])
    
    # Slide 2: Key Modules
    add_slide(prs, "Comprehensive Feature Set", [
        "Restaurant: POS Billing, Table Management, Kitchen Order Ticket (KOT)",
        "Lodging: Reservations, Check-In/Out, Housekeeping, Guest Folio",
        "Common: Inventory Management, Vendors, Advanced Reports",
        "Seamless integration between restaurant and hotel operations."
    ])
    
    # Slide 3: Workflows & Guest Folio
    add_slide(prs, "Integrated Guest Experience", [
        "End-to-end management of the guest journey.",
        "Unified Guest Folio: Combine room charges, restaurant orders, and services.",
        "Efficient billing with itemized invoices.",
        "Real-time status tracking for rooms and kitchen orders."
    ])
    
    # Slide 4: Security & Architecture
    add_slide(prs, "Enterprise-Grade Security", [
        "Multi-Tenant Architecture: Isolated data, roles, and features per client.",
        "Role-Based Access Control (RBAC): Fine-grained permissions.",
        "Roles include: Tenant Admin, Manager, Receptionist, Captain, Kitchen Staff, Housekeeping.",
        "Scalable backend built to handle multiple properties."
    ])
    
    # Slide 5: Brand & Design Philosophy
    add_slide(prs, "Premium Design Aesthetic", [
        "Inspired by luxury boutique hotels and modern systems.",
        "Clean, minimal interfaces that reduce cognitive load.",
        "Color Palette: Deep Classic Navy, Warm Ivory, and Matte Gold.",
        "Typography: 'Lora' for elegance, 'Inter' for usability."
    ])
    
    # Slide 6: The Future
    add_slide(prs, "Future Enhancements", [
        "Multi-Branch Support",
        "Mobile App & QR Menu Ordering",
        "AI Revenue Analytics",
        "Customer Loyalty Program & CRM"
    ])
    
    # Final Slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    title = slide.shapes.title
    title.text = "Thank You"
    title.text_frame.paragraphs[0].font.color.rgb = GOLD
    
    subtitle = slide.placeholders[1]
    subtitle.text = "Ready to elevate your hospitality business?\nContact Us"
    subtitle.text_frame.paragraphs[0].font.color.rgb = IVORY
    
    prs.save("Hoteloraa_Pitch_Deck.pptx")
    print("Presentation generated successfully!")

if __name__ == "__main__":
    create_presentation()
